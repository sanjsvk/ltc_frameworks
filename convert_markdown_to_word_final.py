#!/usr/bin/env python3
"""
Final markdown to Word conversion with robust image and table handling.
"""

import re
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def add_page_numbers(doc):
    """Add page numbers to footer."""
    for section in doc.sections:
        footer = section.footer
        footer.paragraphs[0].text = ""
        run = footer.paragraphs[0].add_run()
        fldChar1 = parse_xml(r'<w:fldChar {} w:fldCharType="begin"/>'.format(nsdecls('w')))
        instrText = parse_xml(r'<w:instrText {} w:space="preserve">PAGE</w:instrText>'.format(nsdecls('w')))
        fldChar2 = parse_xml(r'<w:fldChar {} w:fldCharType="end"/>'.format(nsdecls('w')))
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

def convert_markdown_to_word_final(md_path, output_path):
    """Convert markdown to Word with full image and table support."""
    doc = Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Add page numbers
    add_page_numbers(doc)

    # Read markdown
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')

    i = 0
    table_in_progress = False
    table_lines = []

    while i < len(lines):
        line = lines[i].rstrip()

        # Check if we're starting/in a table
        if line.strip().startswith('| '):
            if not table_in_progress:
                table_in_progress = True
                table_lines = []

            if '---' not in line:  # Skip separator lines
                table_lines.append(line)

            i += 1

            # Check if next line continues table
            if i < len(lines) and not lines[i].strip().startswith('| '):
                # End of table
                if table_lines:
                    _create_word_table(doc, table_lines)
                    table_lines = []
                table_in_progress = False

        # Heading 1
        elif line.startswith('# ') and not line.startswith('## '):
            para = doc.add_paragraph(line[2:].strip(), style='Heading 1')
            para.paragraph_format.line_spacing = 2.0
            i += 1

        # Heading 2
        elif line.startswith('## ') and not line.startswith('### '):
            para = doc.add_paragraph(line[3:].strip(), style='Heading 2')
            para.paragraph_format.line_spacing = 2.0
            i += 1

        # Heading 3
        elif line.startswith('### ') and not line.startswith('#### '):
            para = doc.add_paragraph(line[4:].strip(), style='Heading 3')
            para.paragraph_format.line_spacing = 2.0
            i += 1

        # Heading 4
        elif line.startswith('#### '):
            para = doc.add_paragraph(line[5:].strip(), style='Heading 4')
            para.paragraph_format.line_spacing = 2.0
            i += 1

        # Image
        elif line.strip().startswith('!['):
            match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line.strip())
            if match:
                alt_text, img_path = match.groups()
                # Resolve path from markdown directory
                if not os.path.isabs(img_path):
                    md_dir = os.path.dirname(md_path)
                    full_path = os.path.normpath(os.path.join(md_dir, img_path))
                else:
                    full_path = img_path

                # Try to embed image
                if os.path.exists(full_path):
                    try:
                        para = doc.add_paragraph()
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = para.add_run()
                        run.add_picture(full_path, width=Inches(5.5))
                        para.paragraph_format.space_before = Pt(6)
                        para.paragraph_format.space_after = Pt(6)
                        para.paragraph_format.line_spacing = 2.0
                        print(f"[IMG] Embedded: {alt_text} from {img_path}")
                    except Exception as e:
                        print(f"[ERROR] Failed to embed {img_path}: {e}")
                        para = doc.add_paragraph(f"[Image not embedded: {alt_text}]")
                        para.runs[0].font.italic = True
                else:
                    print(f"[WARN] Image not found: {full_path}")
                    para = doc.add_paragraph(f"[Image not found: {img_path}]")
                    para.runs[0].font.italic = True

            i += 1

        # Empty line
        elif not line.strip():
            i += 1

        # Regular paragraph text
        else:
            if line.strip() and not line.startswith('-'):
                _add_formatted_paragraph(doc, line.strip())
            i += 1

    # Save document
    doc.save(output_path)
    print(f"\n[SUCCESS] Word document created: {output_path}")

def _create_word_table(doc, table_lines):
    """Create a Word table from markdown table lines."""
    if not table_lines:
        return

    # Parse rows
    rows = []
    for line in table_lines:
        cells = [cell.strip() for cell in line.split('|') if cell.strip()]
        rows.append(cells)

    if not rows:
        return

    # Create table
    num_rows = len(rows)
    num_cols = len(rows[0])

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = 'Light Grid Accent 1'

    # Populate cells
    for row_idx, row in enumerate(rows):
        for col_idx, cell_text in enumerate(row):
            cell = table.rows[row_idx].cells[col_idx]
            cell.text = cell_text

            # Format cell
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)
                    run.font.name = 'Calibri'
                paragraph.paragraph_format.line_spacing = 1.5

def _add_formatted_paragraph(doc, text):
    """Add paragraph with inline formatting (bold, italic)."""
    para = doc.add_paragraph()

    # Split by formatting markers
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)

    for part in parts:
        if not part:
            continue

        if part.startswith('**') and part.endswith('**'):
            run = para.add_run(part[2:-2])
            run.font.bold = True
            run.font.size = Pt(12)
            run.font.name = 'Calibri'
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            run = para.add_run(part[1:-1])
            run.font.italic = True
            run.font.size = Pt(12)
            run.font.name = 'Calibri'
        else:
            run = para.add_run(part)
            run.font.size = Pt(12)
            run.font.name = 'Calibri'

    # Format paragraph
    para.paragraph_format.line_spacing = 2.0

if __name__ == '__main__':
    md_file = r'C:\github\ltc\writing\MASTER_DOCUMENT_FINAL.md'
    output_file = r'C:\github\ltc\writing\LTC_Frameworks_JMR.docx'

    print("=" * 70)
    print("MARKDOWN TO WORD CONVERSION (FINAL)")
    print("=" * 70)
    print(f"Input:  {md_file}")
    print(f"Output: {output_file}")
    print("=" * 70)

    convert_markdown_to_word_final(md_file, output_file)
    print("\nDone!")
