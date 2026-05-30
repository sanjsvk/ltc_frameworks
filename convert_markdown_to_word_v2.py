#!/usr/bin/env python3
"""
Enhanced markdown to Word conversion with proper table and image handling.
"""

import re
import os
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def add_page_numbers(doc):
    """Add page numbers to footer."""
    for section in doc.sections:
        footer = section.footer
        footer_para = footer.paragraphs[0]
        footer_para.text = ""

        # Add page number field
        run = footer_para.add_run()
        fldChar1 = parse_xml(r'<w:fldChar {} w:fldCharType="begin"/>'.format(nsdecls('w')))
        instrText = parse_xml(r'<w:instrText {} w:space="preserve">PAGE</w:instrText>'.format(nsdecls('w')))
        fldChar2 = parse_xml(r'<w:fldChar {} w:fldCharType="end"/>'.format(nsdecls('w')))

        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)

        # Center the page numbers
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

def parse_markdown_table(md_table_lines):
    """Parse markdown table and return rows."""
    rows = []
    for line in md_table_lines:
        if line.startswith('| ') and not '-' in line:
            cells = [cell.strip() for cell in line.split('|')[1:-1]]
            rows.append(cells)
    return rows

def markdown_to_word_enhanced(md_path, output_path):
    """Convert markdown to Word with proper table and image support."""
    doc = Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Add page numbers to footer
    add_page_numbers(doc)

    # Read markdown
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')

    i = 0
    while i < len(lines):
        line = lines[i]

        # Heading 1
        if line.startswith('# '):
            para = doc.add_paragraph(line[2:].strip(), style='Heading 1')
            para.paragraph_format.line_spacing = 2.0
            i += 1

        # Heading 2
        elif line.startswith('## '):
            para = doc.add_paragraph(line[3:].strip(), style='Heading 2')
            para.paragraph_format.line_spacing = 2.0
            i += 1

        # Heading 3
        elif line.startswith('### '):
            para = doc.add_paragraph(line[4:].strip(), style='Heading 3')
            para.paragraph_format.line_spacing = 2.0
            i += 1

        # Heading 4
        elif line.startswith('#### '):
            para = doc.add_paragraph(line[5:].strip(), style='Heading 4')
            para.paragraph_format.line_spacing = 2.0
            i += 1

        # Table
        elif line.strip().startswith('| '):
            table_lines = [line]
            i += 1
            # Collect all table lines
            while i < len(lines) and (lines[i].strip().startswith('| ') or lines[i].strip().startswith('|')):
                if not '-' in lines[i]:  # Skip separator lines
                    table_lines.append(lines[i])
                i += 1

            # Parse table
            if len(table_lines) > 1:
                rows = parse_markdown_table(table_lines)
                if rows:
                    # Create table
                    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                    table.style = 'Light Grid Accent 1'

                    # Populate table
                    for row_idx, row in enumerate(rows):
                        for col_idx, cell_text in enumerate(row):
                            cell = table.rows[row_idx].cells[col_idx]
                            cell.text = cell_text
                            # Format cells
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.font.size = Pt(11)
                                    run.font.name = 'Calibri'
                                paragraph.paragraph_format.line_spacing = 1.5

        # Image
        elif line.strip().startswith('!['):
            match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line.strip())
            if match:
                alt_text, img_path = match.groups()
                # Resolve image path
                if not os.path.isabs(img_path):
                    full_path = os.path.join('C:\\github\\ltc', img_path)
                else:
                    full_path = img_path

                if os.path.exists(full_path):
                    try:
                        para = doc.add_paragraph()
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = para.add_run()
                        run.add_picture(full_path, width=Inches(5.5))
                        para.paragraph_format.space_before = Pt(6)
                        para.paragraph_format.space_after = Pt(6)
                        para.paragraph_format.line_spacing = 2.0
                    except Exception as e:
                        # Image embedding failed
                        para = doc.add_paragraph(f"[Image: {alt_text}]")
                        para.runs[0].font.italic = True
                else:
                    para = doc.add_paragraph(f"[Image not found: {img_path}]")
                    para.runs[0].font.italic = True

            i += 1

        # Bold/Italic paragraph
        elif line.strip():
            para = doc.add_paragraph()

            # Process inline formatting
            parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', line.strip())

            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = para.add_run(part[2:-2])
                    run.font.bold = True
                elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
                    run = para.add_run(part[1:-1])
                    run.font.italic = True
                else:
                    para.add_run(part)

            # Format paragraph
            for run in para.runs:
                run.font.size = Pt(12)
                run.font.name = 'Calibri'

            para.paragraph_format.line_spacing = 2.0
            i += 1

        else:
            i += 1

    # Save document
    doc.save(output_path)
    print(f"[OK] Enhanced document created: {output_path}")

if __name__ == '__main__':
    md_file = 'C:\\github\\ltc\\writing\\MASTER_DOCUMENT_FINAL.md'
    output_file = 'C:\\github\\ltc\\writing\\LTC_Frameworks_JMR.docx'

    print("Converting markdown to Word (enhanced)...")
    markdown_to_word_enhanced(md_file, output_file)
    print("Done!")
